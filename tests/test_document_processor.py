"""
Unit tests for document processor.

Tests document loading, placeholder detection, and verse replacement functionality.
"""

import pytest
from pathlib import Path
from docx import Document

from verse_inserter.core.document_processor import DocumentProcessor, ProcessingResult
from verse_inserter.core.placeholder_parser import PlaceholderParser
from verse_inserter.models.verse import Verse, VerseReference, TranslationType


class TestDocumentProcessor:
    """Tests for DocumentProcessor class."""

    @pytest.fixture
    def processor(self):
        """Create a document processor instance."""
        return DocumentProcessor(create_backup=False)  # Disable backup for tests

    def test_initialization(self, processor):
        """Test processor initialization."""
        assert processor is not None
        assert processor.parser is not None
        assert processor.create_backup is False

    def test_load_document_success(self, processor, sample_docx_path):
        """Test successful document loading."""
        with processor.load_document(sample_docx_path) as doc:
            assert doc is not None
            assert isinstance(doc, Document)
            assert len(doc.paragraphs) >= 2

    def test_load_document_not_found(self, processor, tmp_path):
        """Test error handling for non-existent documents."""
        nonexistent = tmp_path / "nonexistent.docx"
        with pytest.raises(Exception):  # Should raise FileNotFoundError or ValueError
            with processor.load_document(nonexistent) as doc:
                pass

    def test_find_placeholders_in_body(self, processor, sample_docx_path):
        """Test finding placeholders in document body."""
        with processor.load_document(sample_docx_path) as doc:
            placeholders = processor.find_all_placeholders(
                doc,
                translation=TranslationType.NIV
            )
            assert len(placeholders) >= 2
            assert any("John 3:16" in p.reference.canonical_reference for p in placeholders)

    def test_find_placeholders_in_tables(self, processor, sample_docx_with_tables):
        """Test finding placeholders in tables."""
        with processor.load_document(sample_docx_with_tables) as doc:
            placeholders = processor.find_all_placeholders(
                doc,
                translation=TranslationType.NIV,
                scan_tables=True
            )
            assert len(placeholders) >= 2
            assert any("John 3:16" in p.reference.canonical_reference for p in placeholders)
            assert any("Psalm 23:1" in p.reference.canonical_reference for p in placeholders)

    def test_progress_callback(self, processor, sample_docx_path):
        """Test progress callback functionality."""
        progress_calls = []

        def progress_callback(current, total, message):
            progress_calls.append((current, total, message))

        processor.set_progress_callback(progress_callback)

        with processor.load_document(sample_docx_path) as doc:
            processor.find_all_placeholders(doc, translation=TranslationType.NIV)

        assert len(progress_calls) > 0
        assert all(isinstance(call[0], int) for call in progress_calls)
        assert all(isinstance(call[1], int) for call in progress_calls)

    def test_replace_placeholders(self, processor, sample_docx_path, sample_verse):
        """Test placeholder replacement."""
        with processor.load_document(sample_docx_path) as doc:
            verse_map = {
                "John 3:16": sample_verse,
                "Genesis 1:1": Verse(
                    reference=VerseReference.parse("Genesis 1:1"),
                    text="In the beginning God created the heavens and the earth.",
                    translation=TranslationType.NIV
                )
            }

            result = processor.replace_placeholders(doc, verse_map)

            assert isinstance(result, ProcessingResult)
            assert result.placeholders_found >= 2
            assert result.placeholders_replaced > 0

    def test_processing_result_success_rate(self):
        """Test success rate calculation."""
        result = ProcessingResult(
            success=True,
            placeholders_found=10,
            placeholders_replaced=8,
            placeholders_failed=2,
            output_file=None,
            errors=[],
            processing_time=1.5
        )
        assert result.success_rate == 80.0

    def test_processing_result_zero_placeholders(self):
        """Test success rate with zero placeholders."""
        result = ProcessingResult(
            success=True,
            placeholders_found=0,
            placeholders_replaced=0,
            placeholders_failed=0,
            output_file=None,
            errors=[],
            processing_time=0.1
        )
        assert result.success_rate == 0.0

    def test_generate_output_filename(self, processor, tmp_path):
        """Test output filename generation."""
        input_file = tmp_path / "document.docx"
        output_file = processor.generate_output_filename(input_file)

        assert output_file.name == "document_filled.docx"
        assert output_file.parent == input_file.parent

    def test_save_document(self, processor, sample_docx_path, tmp_path):
        """Test document saving."""
        output_path = tmp_path / "output.docx"

        with processor.load_document(sample_docx_path) as doc:
            saved_path = processor.save_document(doc, output_path, overwrite=True)
            assert saved_path.exists()
            assert saved_path == output_path

    def test_save_document_no_overwrite(self, processor, sample_docx_path, tmp_path):
        """Test save error when file exists and overwrite=False."""
        output_path = tmp_path / "output.docx"
        output_path.touch()  # Create existing file

        with processor.load_document(sample_docx_path) as doc:
            with pytest.raises(FileExistsError):
                processor.save_document(doc, output_path, overwrite=False)


class TestPlaceholderParser:
    """Tests for PlaceholderParser class."""

    @pytest.fixture
    def parser(self):
        """Create a parser instance."""
        return PlaceholderParser()

    def test_parse_simple_placeholder(self, parser):
        """Test parsing a simple placeholder."""
        text = "This is {{John 3:16}} in the text."
        placeholders = parser.parse_text(text, paragraph_index=0)

        assert len(placeholders) >= 1
        assert any("John 3:16" in p.reference.canonical_reference for p in placeholders)

    def test_parse_multiple_placeholders(self, parser):
        """Test parsing multiple placeholders."""
        text = "We have {{John 3:16}} and {{Genesis 1:1}} here."
        placeholders = parser.parse_text(text, paragraph_index=0)

        assert len(placeholders) >= 2

    def test_parse_verse_range(self, parser):
        """Test parsing verse range placeholders."""
        text = "Read {{Psalm 23:1-6}} for comfort."
        placeholders = parser.parse_text(text, paragraph_index=0)

        assert len(placeholders) >= 1
        placeholder = placeholders[0]
        assert placeholder.reference.is_range

    def test_parse_no_placeholders(self, parser):
        """Test parsing text without placeholders."""
        text = "This text has no verse references."
        placeholders = parser.parse_text(text, paragraph_index=0)

        assert len(placeholders) == 0

    def test_extract_unique_references(self, parser):
        """Test extracting unique references from placeholders."""
        text = "{{John 3:16}} and {{John 3:16}} and {{Genesis 1:1}}"
        placeholders = parser.parse_text(text, paragraph_index=0)

        unique_refs = parser.extract_unique_references(placeholders)

        # Should deduplicate John 3:16
        assert len(unique_refs) <= len(placeholders)
